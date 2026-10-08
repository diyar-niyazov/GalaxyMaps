"""Build the GalaxyMaps round-one and finalist decks (editable PPTX) from real captures.

Run with the python-pptx venv:  /tmp/pptx-venv/bin/python docs/presentation/final/build_decks.py
"""
from pathlib import Path

import cv2
import segno
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

OUT = Path(__file__).resolve().parent
PRESENTATION = OUT.parent
CAP = OUT / "captures"
ASSETS = OUT / "assets"
ASSETS.mkdir(parents=True, exist_ok=True)

BG = RGBColor(0x05, 0x07, 0x0D)
WHITE = RGBColor(0xF4, 0xF1, 0xEA)
MUTED = RGBColor(0xA9, 0xB0, 0xBF)
DIM = RGBColor(0x7C, 0x84, 0x95)
BLUE = RGBColor(0x5B, 0x9B, 0xFF)
ORANGE = RGBColor(0xF5, 0xA6, 0x23)
FONT = "Inter"
URL = "https://www.galaxies.wiki/"
GITHUB = "https://github.com/diyar-niyazov/GalaxyMaps"
TEAM = "TWRD — Towards a new frontier."
MARKS = ASSETS / "marks"

# ---------------------------------------------------------------- image prep


def left_fade(img: Image.Image, width: int, strength: float = 0.94) -> Image.Image:
    """Darken the left side so the title reads; the product image itself is unchanged."""
    w, h = img.size
    mask = Image.new("L", (w, h), 0)
    d = ImageDraw.Draw(mask)
    for x in range(width):
        t = 1 - x / width
        d.line([(x, 0), (x, h)], fill=int(255 * strength * (t ** 1.3)))
    black = Image.new("RGB", (w, h), (5, 7, 13))
    return Image.composite(black, img, mask)


def shifted_hero(src: Path, dx: int, fade: int) -> Image.Image:
    """Hide the quiet-view caption and Restore button (black space), shift the object right."""
    im = Image.open(src).convert("RGB")
    d = ImageDraw.Draw(im)
    d.rectangle([0, 0, 520, 80], fill=(0, 0, 0))
    d.rectangle([1720, 0, 1920, 80], fill=(0, 0, 0))
    canvas = Image.new("RGB", im.size, (0, 0, 0))
    canvas.paste(im, (dx, 0))
    return left_fade(canvas, fade)


def prep_images():
    p = {}
    shifted_hero(CAP / "hero-saturn.png", 330, 1150).save(ASSETS / "cover-saturn.jpg", quality=90)
    shifted_hero(CAP / "hero-earth.png", 420, 1100).save(ASSETS / "closing-earth.jpg", quality=90)
    Image.open(CAP / "route-light.png").convert("RGB").crop((0, 0, 1560, 1080)).save(ASSETS / "route-main.jpg", quality=92)
    Image.open(CAP / "route-voyager.png").convert("RGB").crop((10, 325, 398, 462)).save(ASSETS / "voyager-card.png")
    Image.open(CAP / "route-sources.png").convert("RGB").crop((0, 140, 410, 808)).save(ASSETS / "sources-panel.png")
    Image.open(CAP / "demo2.png").convert("RGB").crop((0, 0, 1420, 1080)).save(ASSETS / "demo2.jpg", quality=92)
    tiles = {
        "explore": ("explore.png", (0, 0, 1100, 619)),
        "directions": ("route-light.png", (0, 0, 1560, 878)),
        "mars": ("route-mars.png", (0, 0, 1560, 878)),
        "compare": ("compare.png", (0, 80, 980, 800)),
        "guide": ("guide.png", (0, 0, 780, 700)),
        "a11y": ("a11y.png", (520, 40, 1400, 535)),
        "tour": ("tour-nebulae.png", (0, 0, 1280, 720)),
        "demo2": ("demo2.png", (0, 40, 820, 780)),
    }
    for name, (src, box) in tiles.items():
        Image.open(CAP / src).convert("RGB").crop(box).resize((1280, 720), Image.LANCZOS).save(ASSETS / f"tile-{name}.jpg", quality=90)

    logo = Image.open(PRESENTATION / "assets" / "logo.png").convert("RGBA")
    px = logo.load()
    for y in range(logo.height):
        for x in range(logo.width):
            r, g, b, a = px[x, y]
            m = min(r, g, b)
            if m > 235:
                px[x, y] = (r, g, b, max(0, int(255 * (255 - m) / 20)))
    logo.crop(logo.getbbox()).save(ASSETS / "logo.png")

    src = ASSETS / "overview-home.jpg"
    if src.exists():
        Image.open(src).convert("RGB").save(ASSETS / "overview-home.jpg", quality=92)

    compose_credits()
    for label, target in (("qr-galaxies-wiki.png", URL), ("qr-github.png", GITHUB)):
        segno.make(target, error="m").save(ASSETS / label, scale=20, border=2, dark="#0B0F1A", light="#FFFFFF")
        decoded, _, _ = cv2.QRCodeDetector().detectAndDecode(cv2.imread(str(ASSETS / label)))
        assert decoded == target, f"{label} decodes to {decoded!r}"
        print("QR verified:", decoded)
    return p


def compose_credits():
    """Right-aligned credit stack: Cursor, Grok, SpaceX + NASA."""
    from PIL import ImageFont
    font = ImageFont.truetype("/usr/share/fonts/inter/InterVariable.ttf", 30)
    cursor = Image.open(MARKS / "cursor.png").convert("RGBA").resize((38, 38), Image.LANCZOS)
    grok = Image.open(MARKS / "grok-white.png").convert("RGBA").resize((36, 36), Image.LANCZOS)
    spacex = Image.open(MARKS / "spacex-white.png").convert("RGBA")
    spacex.thumbnail((168, 28), Image.LANCZOS)
    nasa = Image.open(MARKS / "nasa.png").convert("RGBA")
    nasa.thumbnail((44, 36), Image.LANCZOS)
    ink = (232, 236, 244, 255)

    W, H = 680, 168
    canvas = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    draw = ImageDraw.Draw(canvas)

    def paste_right(items, y, row_h=32):
        x = W - 4
        for item in reversed(items):
            if isinstance(item, str):
                tw = draw.textlength(item, font=font)
                x -= tw
                draw.text((x, y + 3), item, font=font, fill=ink)
                x -= 8
            else:
                x -= item.width
                canvas.paste(item, (int(x), y + (row_h - item.height) // 2), item)
                x -= 8

    paste_right([cursor, "Made with Cursor"], 8)
    paste_right([grok, "Powered by Grok"], 54)
    paste_right(["Data from", spacex, "and", nasa], 100)
    canvas.save(ASSETS / "credits.png")


# ---------------------------------------------------------------- pptx helpers


def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = BG
    return s


def text(slide, x, y, w, h, lines, size=20, color=WHITE, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=1.08, space_after=0):
    """`lines`: list of str or (str, {size,color,bold}) tuples, one paragraph each."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for i, line in enumerate(lines):
        t, o = (line, {}) if isinstance(line, str) else line
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = o.get("spacing", spacing)
        para.space_after = Pt(o.get("after", space_after))
        r = para.add_run()
        r.text = t
        f = r.font
        f.name = FONT
        f.size = Pt(o.get("size", size))
        f.bold = o.get("bold", bold)
        f.color.rgb = o.get("color", color)
    return tb


def picture(slide, path, x, y, w=None, h=None, border=True):
    kw = {}
    if w is not None:
        kw["width"] = Inches(w)
    if h is not None:
        kw["height"] = Inches(h)
    pic = slide.shapes.add_picture(str(path), Inches(x), Inches(y), **kw)
    if border:
        pic.line.color.rgb = RGBColor(0x2A, 0x31, 0x40)
        pic.line.width = Pt(0.75)
    return pic


def notes(slide, body):
    slide.notes_slide.notes_text_frame.text = body.strip()


def eyebrow(slide, label):
    text(slide, 0.7, 0.5, 6, 0.3, [label.upper()], size=12, color=DIM, bold=True)


# ---------------------------------------------------------------- slides

SCRIPTS = {
    "r1": "We’re TWRD — Towards a new frontier. We thought it would be cool to navigate through space. So we built GalaxyMaps: a familiar map that helps people explore space and understand what they’re looking at.",
    "r2": "Search for a destination and get directions. Here’s Earth to Polaris. At light speed, this distance takes about 433 years. At Voyager 1’s speed, it takes millions. You can compare the trips and watch the journey. That’s how we make scale easier to grasp.",
    "r3": "Grok lets you control the map by voice. Ask it to show Saturn, compare planets, or start a tour. It calls validated tools in our Three.js app, while our code calculates the distances and travel times. We built the app with Cursor.",
    "r4": "The same map runs in compatible WebXR headsets. Here’s our Vision Pro demo. You can also explore in a browser, with keyboard controls, spoken descriptions, high contrast, and reduced motion. That gives more students a way in, even without a headset.",
    "r5": "We want a student to ask a question about space and explore the answer. That’s our take on navigation. We’re TWRD, and this is GalaxyMaps.",
    "f1": "We’re TWRD — Towards a new frontier: Teddy, Yijun, Raj, and Diyar. We thought it would be cool to navigate through space. We wanted a way for people to explore space and understand what they’re looking at. This is GalaxyMaps. Navigate the universe.",
    "f2": "Space is vast. Now it’s yours to explore. Whether you're a student or just curious, GalaxyMaps guides you through the stars. Plot a course from Earth to Mars, or venture beyond. Let Grok guide you, or explore with your own hands — in a browser, or inside VR.",
    "f3": "Bring the universe into the classroom. Explore in a browser, or step into the stars with any WebXR headset. Ask questions, navigate by voice, and learn as you go. We designed GalaxyMaps so more people can get in: keyboard and voice, larger text, high contrast, and reduced motion.",
    "f4": "Ask Grok to take you to Saturn and tell you more. Grok is your navigator and your astronomy tutor. It can fly the map and explain what you’re looking at — why Mars is red, how the Sun produces energy, why light can’t escape a black hole.",
    "f5": "Space should be accessible to everyone. Voice, keyboard, mouse, touch, or spatial interaction. You can explore hands-free, and the same universe is there in a browser or a headset.",
    "f6": "We want a student to ask a question about space and explore the answer. That’s our take on navigation. We’re TWRD — Towards a new frontier, and this is GalaxyMaps.",
}


def cover(prs, slot, script):
    s = blank(prs)
    s.shapes.add_picture(str(ASSETS / "cover-saturn.jpg"), 0, 0, prs.slide_width, prs.slide_height)
    picture(s, ASSETS / "logo.png", 0.8, 1.35, h=0.95, border=False)
    text(s, 0.8, 2.45, 7, 1.2, ["GalaxyMaps"], size=72, bold=True)
    text(s, 0.82, 3.65, 7, 0.7, ["Navigate the universe."], size=32, color=WHITE)
    text(s, 0.82, 4.75, 8.5, 0.45, [TEAM], size=22, bold=True, color=BLUE)
    text(s, 0.82, 5.22, 8.5, 0.4, ["Teddy Lampert, Yijun Wang, Raj Iyer, Diyar Niyazov"], size=18, color=WHITE)
    text(s, 0.82, 6.5, 5, 0.6, ["BigRed//Hacks 2026", "Major League Hacking"], size=13, color=MUTED)
    picture(s, ASSETS / "credits.png", 8.45, 5.35, w=4.55, border=False)
    notes(s, f"""{slot}

SAY:
{script}

Visual: genuine GalaxyMaps desktop render of Saturn (quiet view). Saturn texture: Solar System Scope; sky: NASA SVS.""")


def route(prs, slot, script, extra=""):
    s = blank(prs)
    text(s, 0.7, 0.95, 4.7, 0.8, ["Earth to Polaris"], size=40, bold=True)
    text(s, 0.7, 1.85, 4.7, 1.0, ["433 years"], size=58, bold=True, color=BLUE)
    text(s, 0.72, 2.9, 4.6, 0.5, ["at the speed of light"], size=26)
    text(s, 0.72, 4.05, 4.3, 0.4, ["Compare with Voyager 1"], size=18, color=MUTED)
    picture(s, ASSETS / "voyager-card.png", 0.72, 4.5, w=4.0)
    text(s, 0.72, 6.55, 4.6, 0.35, ["Approximate constant-speed comparison"], size=14, color=DIM)
    picture(s, ASSETS / "route-main.jpg", 5.6, 1.05, w=7.25)
    notes(s, f"""{slot}

SAY:
{script}

Demo direction: preload /?route=earth,polaris. Switch Travel mode to Voyager 1 once if reliable; the inset is the genuine Voyager 1 state.
Captured values (map date 2026-10-04): 433 light-years · 133 pc; light speed 433 years (about 5.4 human lifetimes); Voyager 1 (16.9 km/s) 7.7 million years; distance uncertainty ± 6.4 ly (Polaris parallax 7.54 ± 0.11 mas).
Straight-line cruise at constant speed; positions frozen at the map date.{extra}""")


def overview(prs, slot, script):
    s = blank(prs)
    picture(s, ASSETS / "overview-home.jpg", 5.25, 0.85, w=7.7)
    text(s, 0.55, 0.65, 4.55, 1.85, ["Space is vast.", "Now it’s yours", "to explore."], size=36, bold=True, spacing=1.0)
    text(s, 0.57, 2.65, 4.5, 1.1, ["Whether you're a student or just curious, GalaxyMaps guides you through the stars and helps you discover awesome places."], size=16, color=WHITE, spacing=1.15)
    text(s, 0.57, 3.9, 4.5, 1.55, [
        "Plot your course.",
        "Travel from Earth to Mars, or venture beyond.",
        "Let Grok guide you, or take exploration into your own hands.",
    ], size=16, color=WHITE, space_after=8, spacing=1.12)
    text(s, 0.57, 5.7, 4.5, 0.75, ["Explore in your browser.", "Step inside with VR."], size=18, bold=True, color=BLUE, spacing=1.1)
    notes(s, f"""{slot}

SAY:
{script}

Visual: genuine GalaxyMaps home view supplied by the team (Explore panel, Earth locked, Earth → Mars journey).""")


def saturn_vr(prs, slot, script):
    s = blank(prs)
    text(s, 0.5, 0.5, 4.55, 1.9, ["“Grok, take me", "to Saturn and", "tell me more.”"], size=30, bold=True, spacing=1.02)
    text(s, 0.52, 2.5, 4.55, 0.7, ["Your navigator. Your astronomy tutor. Your companion through the cosmos."], size=16, color=BLUE, spacing=1.12)
    text(s, 0.52, 3.25, 4.55, 3.55, [
        "Grok combines scientific knowledge with voice navigation, turning every destination into a hands-on lesson.",
        "Ask why Mars is red, how the Sun produces energy, or why light can’t escape a black hole.",
        "With immersive visuals and Grok beside them, students can move through space, examine worlds up close, and connect each explanation to what they see.",
    ], size=16, color=WHITE, space_after=11, spacing=1.14)
    picture(s, ASSETS / "02_VisionPro_Immersive_Saturn_Grok.jpeg", 5.3, 0.85, w=7.55)
    text(s, 5.3, 5.52, 7.55, 0.35, ["Apple Vision Pro · immersive view with the Grok mic HUD"], size=13, color=MUTED)
    notes(s, f"""{slot}

SAY:
{script}

Visual: genuine Apple Vision Pro capture of the immersive WebXR scene: Saturn, the Grok mic HUD and the in-world Saturn card. It shows the interface, not a completed spoken command; don't describe it as a live response.
Grok can act through validated app tools and may use outside knowledge to explain objects. Do not claim a completed spoken command from this still.""")


def classroom(prs, slot, script):
    s = blank(prs)
    picture(s, ASSETS / "01_VisionPro_Earth_Hand.jpeg", 0.6, 1.15, w=8.3)
    text(s, 0.6, 5.95, 8.3, 0.35, ["Apple Vision Pro demo · GalaxyMaps in the spatial browser"], size=13, color=MUTED)
    text(s, 9.15, 0.65, 3.85, 1.4, ["Bring the universe", "into the classroom."], size=28, bold=True, spacing=1.0)
    text(s, 9.17, 2.25, 3.8, 0.85, ["Explore in a browser, or step into the stars with any WebXR-compatible VR headset."], size=15, color=WHITE, spacing=1.12)
    text(s, 9.17, 3.2, 3.8, 0.35, ["Let curiosity lead."], size=16, bold=True, color=BLUE)
    text(s, 9.17, 3.55, 3.8, 0.7, ["Ask questions. Navigate by voice. Learn as you explore."], size=15, color=WHITE, spacing=1.12)
    text(s, 9.17, 4.4, 3.8, 0.4, ["Making space accessible to all."], size=16, bold=True, color=BLUE)
    text(s, 9.17, 4.85, 3.8, 1.2, ["Both keyboard and voice controlled navigation, larger text, high contrast, and reduced motion."], size=15, color=WHITE, spacing=1.12)
    notes(s, f"""{slot}

SAY:
{script}

Visual: genuine Apple Vision Pro capture: GalaxyMaps (Earth) in a spatial browser window, with the room and a pinch gesture visible. It is a still; don't describe it as showing motion.
Accessibility options named here exist in the app's Accessibility panel: keyboard navigation, Describe current view (spoken descriptions), larger text, high contrast, reduced motion. Voice control uses Grok. The app aims toward WCAG 2.2 AA and states its known limitations; it does not claim full conformance.""")


def closing(prs, slot, script):
    s = blank(prs)
    s.shapes.add_picture(str(ASSETS / "closing-earth.jpg"), 0, 0, prs.slide_width, prs.slide_height)
    text(s, 0.8, 1.0, 5.6, 2.0, ["Where would", "you go?"], size=54, bold=True, spacing=1.0)
    picture(s, ASSETS / "qr-galaxies-wiki.png", 0.8, 3.25, w=1.7, border=False)
    picture(s, ASSETS / "qr-github.png", 2.7, 3.25, w=1.7, border=False)
    text(s, 0.8, 5.1, 1.7, 0.35, ["galaxies.wiki"], size=13, bold=True, align=PP_ALIGN.CENTER)
    text(s, 2.7, 5.1, 1.7, 0.35, ["GitHub"], size=13, bold=True, align=PP_ALIGN.CENTER)
    text(s, 0.8, 5.55, 5.6, 0.4, [TEAM], size=18, bold=True, color=BLUE)
    notes(s, f"""{slot}

SAY:
{script}

QR codes verified: {URL} and {GITHUB}. Leave this slide up during Q&A.""")


def under_the_map(prs, slot, script):
    s = blank(prs)
    text(s, 0.7, 0.95, 6.6, 0.8, ["Under the map"], size=40, bold=True)
    text(s, 0.72, 1.85, 6.4, 1.0, ["Astronomy data and route math power every trip."], size=22, spacing=1.15)
    labels = ["Precise 3D positions", "Camera-relative rendering", "Separate travel models", "Validated Grok tools"]
    for i, lab in enumerate(labels):
        y = 3.15 + i * 0.68
        sq = s.shapes.add_shape(1, Inches(0.74), Inches(y + 0.1), Inches(0.14), Inches(0.14))
        sq.fill.solid()
        sq.fill.fore_color.rgb = BLUE
        sq.line.fill.background()
        text(s, 1.1, y, 5.8, 0.45, [lab], size=22)
    text(s, 0.72, 6.35, 6, 0.35, ["React, TypeScript, Three.js, WebXR"], size=14, color=DIM)
    picture(s, ASSETS / "sources-panel.png", 7.9, 0.75, h=6.0)
    notes(s, f"""{slot}

SAY:
{script}

Visual: the genuine Earth → Polaris "Sources and model details" panel (constant-speed straight line, positions frozen at the map date, SI definition of c, JPL Horizons position for Earth, Polaris parallax 7.54 ± 0.11 mas).
Verified in the repository: catalog positions are float64 ICRF kilometres, projected relative to the camera every frame (docs/ARCHITECTURE.md); route math lives in TypeScript separate from rendering; the straight-line benchmark and the idealized Hohmann transfer are separate models; Grok can act only through validated app tools. Don't quote frame rates or performance numbers.""")


def demo2(prs, slot, script):
    s = blank(prs)
    text(s, 0.7, 1.0, 4.4, 1.5, ["Demo-2: launch to homecoming"], size=36, bold=True, spacing=1.0)
    text(s, 0.72, 2.65, 4.3, 0.9, ["Five chapters from the real mission"], size=22, color=BLUE, spacing=1.1)
    text(s, 0.72, 3.85, 4.3, 1.2, ["NASA photographs and source links", "SpaceX spacecraft profiles"], size=18, space_after=10)
    picture(s, ASSETS / "demo2.jpg", 5.35, 0.9, h=5.7)
    notes(s, f"""{slot}

SAY:
{script}

Visual: the genuine Demo-2 mission story (chapter 2 of 5, "Liftoff from Kennedy"), with the NASA/Bill Ingalls photograph credit and NASA source link preserved. The story says the map shows destination context and does not reconstruct a flight trajectory; don't call it a trajectory replay.""")


def vr_access(prs, slot, script):
    s = blank(prs)
    picture(s, ASSETS / "visionpro-andromeda.png", 0.5, 0.75, w=8.15)
    text(s, 0.5, 6.15, 8.15, 0.4, ["GalaxyMaps on Apple Vision Pro · spatial exploration with natural interaction"], size=13, color=MUTED)
    text(s, 9.15, 0.65, 3.85, 1.45, ["Space should be", "accessible to", "everyone."], size=24, bold=True, spacing=1.0)
    text(s, 9.17, 2.25, 3.8, 0.4, ["Voice-controlled navigation"], size=15, bold=True, color=BLUE)
    text(s, 9.17, 2.65, 3.8, 0.55, ["Explore and ask questions hands-free."], size=14, color=WHITE, spacing=1.1)
    text(s, 9.17, 3.3, 3.8, 0.4, ["Accessible by design"], size=15, bold=True, color=BLUE)
    text(s, 9.17, 3.7, 3.8, 0.85, ["Keyboard controls, larger text, high contrast, and reduced motion."], size=14, color=WHITE, spacing=1.1)
    text(s, 9.17, 4.65, 3.8, 0.4, ["One universe. Many ways to explore."], size=15, bold=True, color=BLUE)
    text(s, 9.17, 5.05, 3.8, 0.85, ["Browser, VR, voice, touch, mouse, or keyboard."], size=14, color=WHITE, spacing=1.1)
    notes(s, f"""{slot}

SAY:
{script}

Visual: genuine Apple Vision Pro capture of Andromeda with the destination card and Grok HUD. Designed with accessibility in mind; do not claim audited ADA or WCAG conformance.""")


FEATURES = [
    ("explore", "Explore", "Search, Surprise me, saved places"),
    ("directions", "Earth to Polaris", "Light speed · Voyager 1 · 433 years"),
    ("mars", "Earth to Mars", "Idealized Hohmann · 8.5 months"),
    ("compare", "Compare sizes", "True scale · Earth and Jupiter"),
    ("guide", "Grok", "Voice or typed commands"),
    ("a11y", "Accessibility", "Keyboard, speech, contrast, motion"),
    ("tour", "Guided tours", "Beautiful nebulae · sourced photos"),
    ("demo2", "Demo-2", "Five NASA-sourced chapters"),
]


def features_backup(prs):
    s = blank(prs)
    text(s, 0.42, 0.18, 12.4, 0.42, ["More of GalaxyMaps"], size=26, bold=True)
    w, gap_x, gap_y = 3.02, 0.16, 0.42
    h = w * 9 / 16
    x0 = (13.333 - (4 * w + 3 * gap_x)) / 2
    y0 = 0.72
    for i, (key, title, caption) in enumerate(FEATURES):
        x = x0 + (i % 4) * (w + gap_x)
        y = y0 + (i // 4) * (h + gap_y)
        picture(s, ASSETS / f"tile-{key}.jpg", x, y, w=w)
        text(s, x, y + h + 0.04, w, 0.36, [(title, {"size": 13, "bold": True}), (caption, {"size": 11, "color": MUTED})], spacing=1.0)
    text(s, 0.42, 7.12, 12.5, 0.28, ["Cited catalogs · app-calculated distances · React, TypeScript, Three.js, WebXR · keys stay on the server · built with Cursor · next: a classroom pilot"], size=11, color=DIM)
    notes(s, """Backup for Q&A, not timed. Replaces the six text inventory slides. Every tile is a genuine capture of the current app.

Explore: search, categories, Surprise me, saved places, Realistic and Atlas layers, shareable views, Play time.
Earth to Polaris: multi-stop routes, light speed and Voyager 1, journey playback, suggested detours. Captured: 433 years at c, 7.7 million years at Voyager 1.
Earth to Mars: idealized Hohmann transfer (~8.5 months), separate from the constant-speed comparison. Assumptions and uncertainty are in Sources and model details.
Compare sizes: Earth/Jupiter (also Earth/Sun, Sun/Sirius A, Earth/Moon).
Grok: voice and typed Mission Control through validated app tools; dictation and spoken replies. Needs a microphone and a configured server.
Accessibility: keyboard, Describe current view, larger text, high contrast, reduced motion. Aims toward WCAG 2.2 AA; does not claim full conformance.
Guided tours: Beautiful nebulae and other self-paced tours with sourced photographs.
Demo-2: five-chapter educational story with NASA photographs; Falcon 9, Dragon, and Starship profiles. No SpaceX API or telemetry.
Classroom next step (proposed, not yet run): predict travel time, change the mode, compare, explain.""")


def build(name, main):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    for fn in main:
        fn(prs)
    under_the_map(prs, "Backup", "Positions stay high-precision and render relative to the camera. Route math is separate from the display. Constant-speed comparisons and idealized orbital transfers are labeled differently.")
    demo2(prs, "Backup", "Our Demo-2 story follows the mission through five chapters, with NASA photographs and source links. Falcon 9, Dragon, and Starship profiles sit alongside it.")
    features_backup(prs)
    path = OUT / f"{name}.pptx"
    prs.save(path)
    print("wrote", path, len(prs.slides), "slides")


def main():
    prep_images()
    S = SCRIPTS
    build("GalaxyMaps-Round-One", [
        lambda p: cover(p, "Slide 1 · 0:00–0:14", S["r1"]),
        lambda p: overview(p, "Slide 2 · 0:14–0:42", S["f2"]),
        lambda p: classroom(p, "Slide 3 · 0:42–1:12", S["f3"]),
        lambda p: saturn_vr(p, "Slide 4 · 1:12–1:37", S["f4"]),
        lambda p: closing(p, "Slide 5 · 1:37–1:50 (buffer to 2:00)", S["r5"]),
    ])
    build("GalaxyMaps-Finalist", [
        lambda p: cover(p, "Slide 1 · 0:00–0:25", S["f1"]),
        lambda p: overview(p, "Slide 2 · 0:25–1:05", S["f2"]),
        lambda p: classroom(p, "Slide 3 · 1:05–1:50", S["f3"]),
        lambda p: saturn_vr(p, "Slide 4 · 1:50–2:40", S["f4"]),
        lambda p: vr_access(p, "Slide 5 · 2:40–3:20", S["f5"]),
        lambda p: closing(p, "Slide 6 · 3:20–3:45 (buffer to 4:00)", S["f6"]),
    ])
    for k in ("r", "f"):
        words = sum(len(v.split()) for kk, v in S.items() if kk.startswith(k))
        print(k, "script words:", words)


if __name__ == "__main__":
    main()
